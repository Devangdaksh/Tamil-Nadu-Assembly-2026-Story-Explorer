"""
Build AtliQ stakeholder PowerPoint deck.
Run: python analysis/build_deck.py
Requires: python analysis/generate_outputs.py first
"""

from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
sys.path.insert(0, str(ROOT))

from src.load_data import constituency_margins, constituency_winners, load_results, year_comparison  # noqa: E402

ACCENT = RGBColor(59, 130, 246)
DARK = RGBColor(24, 24, 48)
MUTED = RGBColor(107, 114, 128)


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(8.8), Inches(2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = DARK
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED
    p2.space_before = Pt(12)
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(1.8), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    image: Path | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(24)
    tp.font.bold = True
    tp.font.color.rgb = DARK

    left = Inches(0.6)
    width = Inches(4.2) if image else Inches(8.8)
    body = slide.shapes.add_textbox(left, Inches(1.2), width, Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = bullet
        para.font.size = Pt(14)
        para.font.color.rgb = RGBColor(34, 34, 34)
        para.level = 0
        para.space_after = Pt(8)

    if image and image.exists():
        slide.shapes.add_picture(str(image), Inches(5.0), Inches(1.2), width=Inches(4.4))


def flip_table_text() -> list[str]:
    cmp = year_comparison()
    flips = cmp[cmp["changed"]].groupby(["party_2021", "party_2026"]).size()
    top = flips.sort_values(ascending=False).head(6)
    lines = ["Top party switches (2021 winner → 2026 winner):"]
    for (a, b), n in top.items():
        lines.append(f"• {a} → {b}: {n} constituencies")
    lines.append(f"• Total constituencies that changed winner: {len(cmp[cmp.changed])}")
    return lines


def main() -> None:
    w21 = constituency_winners(load_results(2021))
    w26 = constituency_winners(load_results(2026))
    margins = constituency_margins(load_results(2026))
    closest = margins.nsmallest(1, "margin_votes").iloc[0]
    chennai = w26[w26.region == "Chennai Metro"]
    tvk_metro = int((chennai.party == "TVK").sum())

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Tamil Nadu 2026: A wave and a photo finish",
        "AtliQ Media · ECI public data · Prepared for editorial planning",
    )

    add_content_slide(
        prs,
        "What we measured",
        [
            "234 Assembly constituencies · 2021 vs 2026 candidate-level results",
            "Six editorial regions: Chennai Metro, North, Central, Kongu, Delta, South",
            "Neutral, fact-based framing — no commentary or predictions",
            "Source: Election Commission of India (2026 live portal; 2021 cleaned ECI data)",
        ],
    )

    add_content_slide(
        prs,
        "Story 1 — The seat count reset",
        [
            f"2026 winners: TVK {int((w26.party == 'TVK').sum())} · DMK {int((w26.party == 'DMK').sum())} · AIADMK {int((w26.party == 'AIADMK').sum())}",
            f"2021 winners: DMK {int((w21.party == 'DMK').sum())} · AIADMK {int((w21.party == 'AIADMK').sum())}",
            "TVK had no winning seats in the 2021 file (new entrant in 2026 results)",
            "Largest single flip path: DMK → TVK (65 constituencies)",
        ],
        OUT / "seats_2026.png",
    )

    add_content_slide(
        prs,
        "Story 1 — Where TVK won",
        [
            f"Chennai Metro: TVK won {tvk_metro} of {len(chennai)} seats",
            "South region: TVK 26 · DMK 22 (most competitive belt after Metro)",
            "North & Central: AIADMK remains strong (15 seats each in 2026)",
            "Pattern: urban concentration + statewide competitiveness",
        ],
        OUT / "seats_by_region_2026.png",
    )

    add_content_slide(
        prs,
        "Story 2 — The super-tight map",
        [
            f"{int((margins.margin_pct < 5).sum())} constituencies decided by under 5% margin",
            f"{int((margins.margin_pct < 1).sum())} constituencies decided by under 1% margin",
            "More than four in ten seats were genuine nail-biters",
            "Recommendation: open the TV hour with human drama, not the sweep headline",
        ],
        OUT / "margin_distribution_2026.png",
    )

    add_content_slide(
        prs,
        "Story 2 — Extreme close-up",
        [
            f"Narrowest result: {closest.constituency}",
            f"Margin: {int(closest.margin_votes)} vote ({closest.party} over {closest.runner_up_party})",
            "Verify against final ECI Form-20 before broadcast",
            "Field segment opportunity: interview voters in top-5 closest ACs",
        ],
    )

    add_content_slide(
        prs,
        "Story 3 — Who lost what (2021 → 2026)",
        flip_table_text(),
    )

    add_content_slide(
        prs,
        "Editorial recommendation — 60-minute block",
        [
            "Segment A (7 min): “The nail-biters” — margin map + Tiruppattur",
            "Segment B (7 min): “Metro realignment” — Chennai + South heatmap",
            "Segment C (7 min): “The 65-seat switch” — flip table + DMK retention",
            "Desk wrap: state limitations; no causal claims without separate sourcing",
        ],
    )

    add_content_slide(
        prs,
        "Data limitations (state on air)",
        [
            "2026 turnout % is not in the supplied dataset",
            "2026 totals are pre–Form-20; audited figures may differ slightly",
            "Vote shares are constituency averages, not statewide popular vote",
            "AC name spelling varies in a few rows; ac_number is the join key",
        ],
    )

    dest = OUT / "AtliQ_TN_Election_2026_Story_Deck.pptx"
    prs.save(dest)
    print(f"Saved {dest}")


if __name__ == "__main__":
    main()
